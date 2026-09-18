"""Independent, read-only second-pass QA of the frozen D5 to D6 release.

Run after the offline pipeline and workbook builder. This deliberately reads
their saved artifacts rather than importing the D6 calculation modules.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
D5 = ROOT / "D5"
D6 = ROOT / "D6_cost_analysis"
OUT = D6 / "outputs"
XLSX = D6 / "submission" / "PE6201_D6_Cost_Analysis_Teacher_Submission_FIXED.xlsx"

FROZEN_SHA256 = {
    "D5/FINAL_5PLUS1_QA.md": "BE9B0D1228BB466C66C93DCC4EB7276A59C104460765488F6E1A0DA0F9F36AA3",
    "D5/D5_NORMALIZED_SCORE_CHANGES.csv": "B747D93283CB7394AA8F1E72CB620AD29CE43ABC962568EB24A94AB7A3E5086F",
    "D5/SCORING_NORMALIZATION_AUDIT.md": "378E8B5B1F9E15451DEEB7A892E86FA74F0197FE68D1764FF8D2AFC20C996EBC",
    "D5/outputs/D5_MINIMAL_GITHUB_PACKAGE/SELECTED_5PLUS1_INVENTORY.csv": "824B0A69E15FD01E3D6EF7D4BE44164490D52CB69A907ACEA219E3CF97513E2F",
    "D5/outputs/D5_MINIMAL_GITHUB_PACKAGE/results/live/SELECTED_FINAL_INDEX.json": "4EE79A84FF3C02AB99DCD69A7685D0C6D9044C854D34637F20EF1D6B583C4FC8",
    "D5/outputs/D5_MINIMAL_GITHUB_PACKAGE/D5_COMPARISON.md": "EDD927B78DFB1D75B7F5795ECBC0503929C2ED44744400B914DDED32BBCFE249",
    "D5/outputs/D5_MINIMAL_GITHUB_PACKAGE/D5_COST_RECONCILIATION.md": "7A4AB7DA3C2C2B3A28F8DD1B91569D0FE5DF57C46BE9ED697C1F8DB511716B9A",
}
EXPECTED = {
    "openai_gpt4o_mini_v2": (52, 17, 35, 0, 14, 4, 0, 29644.075883888887),
    "qwen3_30b_v2": (52, 28, 14, 7, 6, 2, 7, 14672.840641666666),
    "mistral_small_3_2_v2": (52, 37, 0, 4, 9, 0, 4, 8258.426726666665),
    "gemini2_5_flash_v2": (52, 39, 6, 0, 12, 6, 0, 8266.09277),
    "claude_opus5_frontier_negative_v2": (18, 10, 8, 0, 10, 8, 0, None),
    "qwen3_30b_v1": (52, 20, 26, 3, None, None, None, None),
}
V2_ORDER = list(EXPECTED)[:5]
LEGACY = [
    "outputs/missing_cost_sensitivity.csv",
    "outputs/pareto_frontier.csv",
    "preflight/d6_readiness_matrix.csv",
    "preflight/experiment_integrity.md",
    "preflight/phase_a_summary.md",
    "preflight/repository_audit.md",
    "preflight/scripts/audit_phase_a.py",
]
STALE = ("llama", "llama3_3_70b", "51/52", "311/312", "312 formal", "pricing_mismatch")


def rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def dec(value: object) -> Decimal:
    return Decimal(str(value))


def near(actual: object, expected: object, tolerance: float = 1e-8) -> None:
    assert math.isclose(float(actual), float(expected), abs_tol=tolerance, rel_tol=0), (actual, expected)


def main() -> None:
    checks = 0

    def check(condition: bool, message: str) -> None:
        nonlocal checks
        checks += 1
        assert condition, message

    for relative, expected_digest in FROZEN_SHA256.items():
        digest = hashlib.sha256((ROOT / relative).read_bytes()).hexdigest().upper()
        check(digest == expected_digest, f"Frozen D5 hash drift: {relative}: {digest}")
    print(f"PASS: frozen D5 SHA-256 unchanged ({len(FROZEN_SHA256)}/{len(FROZEN_SHA256)} key files)")

    index = json.loads((D5 / "outputs/D5_MINIMAL_GITHUB_PACKAGE/results/live/SELECTED_FINAL_INDEX.json").read_text(encoding="utf-8"))
    experiments = index["selected_v2"] + index["prompt_control"]
    check([e["experiment_id"] for e in experiments] == list(EXPECTED), "Selected D5 5+1 inventory/order")
    source = rows(D5 / "D5_NORMALIZED_SCORE_CHANGES.csv")
    actual = rows(OUT / "formal_runs_normalized.csv")
    check(len(source) == len(actual) == 278, "278 selected source/output rows")
    check(sum(r["changed"] == "True" for r in source) == 0, "D5 zero normalized score changes")
    check(len({(r["experiment_id"], r["case_id"], r["trial"]) for r in source}) == 278, "D5 unique keys")
    src_by_key = {(r["experiment_id"], r["case_id"], r["trial"]): r for r in source}
    out_by_key = {(r["experiment_id"], r["case_id"], r["trial"]): r for r in actual}
    check(src_by_key.keys() == out_by_key.keys(), "D5/D6 selected keys")
    for key, src in src_by_key.items():
        dst = out_by_key[key]
        check(src["normalized_passed"] == dst["passed"], f"Pass value {key}")
        for field in ("input_tokens", "output_tokens", "cached_input_tokens", "reasoning_tokens", "provider_cost_usd"):
            check(src[field] == dst[field], f"Raw {field} value {key}")
    check(all(r["provider_cost_usd"] for r in actual), "278/278 provider cost coverage")
    check(not any("llama" in (r["model"] + r["experiment_id"]).lower() for r in actual), "No selected Llama")
    spend = sum((dec(r["provider_cost_usd"]) for r in source), Decimal(0))
    check(spend == dec("2.13502002"), f"D5 scored spend {spend}")

    by_experiment = defaultdict(list)
    for row in source:
        by_experiment[row["experiment_id"]].append(row)
    common_keys = None
    for experiment, (n, passed, invalid, unsafe, neg_pass, neg_invalid, neg_unsafe, _) in EXPECTED.items():
        group = by_experiment[experiment]
        check(len(group) == n, f"{experiment} source denominator")
        check(sum(r["normalized_passed"] == "True" for r in group) == passed, f"{experiment} source passes")
        check(sum(r["status"] == "invalid_model_output" for r in group) == invalid, f"{experiment} source invalid")
        check(sum(r["unsafe_book_slot_attempt"] == "True" for r in group if r["negative_case"] == "True") == unsafe, f"{experiment} source unsafe")
        if experiment in V2_ORDER:
            negative = [r for r in group if r["negative_case"] == "True"]
            keys = {(r["case_id"], r["trial"]) for r in negative}
            check(len(negative) == len(keys) == 18, f"{experiment} negative denominator/keys")
            check({r["case_id"] for r in negative} == {f"REF-{n}" for n in range(6060, 6066)}, f"{experiment} negative cases")
            if common_keys is None:
                common_keys = keys
            check(keys == common_keys, f"{experiment} matched negative keys")
            check(sum(r["normalized_passed"] == "True" for r in negative) == neg_pass, f"{experiment} negative passes")
            check(sum(r["status"] == "invalid_model_output" for r in negative) == neg_invalid, f"{experiment} negative invalid")
            check(sum(r["unsafe_book_slot_attempt"] == "True" for r in negative) == neg_unsafe, f"{experiment} negative unsafe")
    check(sum(r["status"] == "invalid_model_output" for r in source) == 89, "89 invalid source rows")
    claude_source = by_experiment["claude_opus5_frontier_negative_v2"]
    check(sum(dec(r["provider_cost_usd"]) for r in claude_source) == dec("1.644405"), "Claude source spend")
    for field, expected in (("input_tokens", 272906), ("output_tokens", 11195), ("reasoning_tokens", 502), ("cached_input_tokens", 0)):
        check(sum(int(r[field]) for r in claude_source) == expected, f"Claude {field}")
    print("PASS: independent D5 row-level selection, scores, usage, safety, cost and matched negative keys")

    summary = {r["experiment_id"]: r for r in rows(OUT / "model_cost_summary.csv")}
    negative = {r["experiment_id"]: r for r in rows(OUT / "common_negative_cost_context.csv")}
    check(set(summary) == set(EXPECTED), "Six model summary rows")
    check(set(negative) == set(V2_ORDER), "Five negative model rows")
    report = json.loads((OUT / "cost_report.json").read_text(encoding="utf-8"))
    report_summary = {r["experiment_id"]: r for r in report["model_summary"]}
    md = (OUT / "cost_report.md").read_text(encoding="utf-8")
    for experiment, (n, passed, invalid, unsafe, neg_pass, neg_invalid, neg_unsafe, case_monthly) in EXPECTED.items():
        s = summary[experiment]
        check(int(s["runs"]) == n and int(s["passed_runs"]) == passed, f"{experiment} CSV pass")
        check(int(s["invalid_model_output_count"]) == invalid and int(s["unsafe_action_count"]) == unsafe, f"{experiment} CSV invalid/unsafe")
        check(s["provider_cost_coverage"] == f"{n}/{n}", f"{experiment} CSV cost coverage")
        check(report_summary[experiment]["passed_runs"] == passed, f"{experiment} report JSON pass")
        if experiment in V2_ORDER:
            v = negative[experiment]
            check((int(v["runs"]), int(v["passed"]), int(v["invalid_model_output"]), int(v["unsafe_booking_attempts"])) == (18, neg_pass, neg_invalid, neg_unsafe), f"{experiment} negative CSV")
            check(f"{neg_pass}/18" in md, f"{experiment} negative Markdown numerator")
        if case_monthly is not None:
            near(s["case_balanced_monthly_cost"], case_monthly)
            near(report_summary[experiment]["case_balanced_monthly_cost"], case_monthly)
            check(f"{passed}/52" in md, f"{experiment} full Markdown numerator")
        elif experiment == "claude_opus5_frontier_negative_v2":
            check(s["case_balanced_monthly_cost"] == "" and s["trial_weighted_monthly_cost"] == "", f"{experiment} no general monthly")
    check(summary["claude_opus5_frontier_negative_v2"]["scope"] == "FRONTIER_NEGATIVE_ONLY", "Claude scope")
    check(summary["claude_opus5_frontier_negative_v2"]["overall_pass_rate_status"] == "N/A_NEGATIVE_ONLY_SCOPE", "Claude overall N/A")
    check(dec(summary["claude_opus5_frontier_negative_v2"]["provider_spend_usd"]) == dec("1.644405"), "Claude CSV spend")
    for filename in ("sensitivity_analysis.csv", "break_even_matrix.csv", "break_even_pairs.csv", "bootstrap_uncertainty.csv", "pareto_analysis.csv"):
        content = (OUT / filename).read_text(encoding="utf-8").lower()
        check("claude" not in content and "llama" not in content, f"No Claude/Llama general economics in {filename}")
    evaluation = {r["metric"]: dec(r["value_usd"]) for r in rows(OUT / "evaluation_spend_reconciliation.csv")}
    for key, value in (("selected_scored_run_spend", "2.13502002"), ("extra_charged_mistral_provider_error", "0.00031845"), ("recorded_selected_model_charges_including_error", "2.13533847")):
        check(evaluation[key] == dec(value), f"Spend CSV {key}")
        check(dec(report["evaluation_spend"][key]) == dec(value), f"Spend JSON {key}")
        check(value in md, f"Spend Markdown {key}")
    qwen = {r["weighting_view"]: r for r in rows(OUT / "qwen_prompt_ablation.csv")}
    check(set(qwen) == {"trial_weighted", "case_balanced"}, "Qwen weighting views")
    near(qwen["trial_weighted"]["delta_pass_rate_pp"], 15.38461538461538)
    check("+15.38 pp" in md, "Qwen Markdown delta")
    check("2/18" in md and "6/18" in md, "Correct Qwen/Gemini negative invalid Markdown")
    print("PASS: D6 CSV/JSON/Markdown economics, scope and spend consistency")

    expected_sheets = ["Executive_Summary", "Model_Cost", "Prompt_Impact", "Evaluation_Spend", "Sensitivity", "Break_Even", "Uncertainty", "Cost_Levers", "Safety_Reliability", "Cost_Guardrails", "Methodology_Limits"]
    values = load_workbook(XLSX, data_only=True, read_only=False)
    formulas = load_workbook(XLSX, data_only=False, read_only=False)
    check(values.sheetnames == expected_sheets, "Workbook sheets/order")
    errors = []
    for sheet in values:
        check(sheet.sheet_state == "visible", f"Visible sheet {sheet.title}")
        for row in sheet:
            for cell in row:
                if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value in {"#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?"}):
                    errors.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
                if isinstance(cell.value, str):
                    check("llama" not in cell.value.lower(), f"No workbook Llama {sheet.title}!{cell.coordinate}")
        check(not any(d.hidden for d in sheet.row_dimensions.values()), f"No hidden rows in {sheet.title}")
    check(not errors, f"Workbook formula errors: {errors}")
    model = values["Model_Cost"]
    executive = values["Executive_Summary"]
    for offset, experiment in enumerate(V2_ORDER[:4]):
        s = summary[experiment]
        r = 11 + offset
        check(model[f"B{r}"].value == int(s["passed_runs"]) and model[f"C{r}"].value == 52, f"Workbook {experiment} denominator")
        near(model[f"D{r}"].value, s["trial_weighted_pass_rate"])
        near(model[f"I{r}"].value, s["trial_weighted_monthly_cost"], tolerance=1e-5)
        near(model[f"G{19+offset}"].value, s["case_balanced_monthly_cost"], tolerance=1e-5)
        check(executive[f"B{12+offset}"].value == f"{s['passed_runs']}/52", f"Executive {experiment} numerator")
        near(executive[f"F{12+offset}"].value, s["trial_weighted_monthly_cost"], tolerance=1e-5)
        check(formulas["Model_Cost"][f"I{r}"].data_type == "f", f"Workbook live formula {experiment}")
    check(model["B27"].value == 10 and model["C27"].value == 18, "Workbook Claude negative denominator")
    near(model["G27"].value, 1.644405)
    check(model["I27"].value is None, "Workbook Claude no monthly figure")
    near(executive["B20"].value, 2.13502002)
    near(executive["B21"].value, 0.00031845)
    check("Claude" in str(executive["A25"].value) and "Frontier" in str(executive["A27"].value), "Executive Claude/tier note")
    check(all(values["Safety_Reliability"][f"B{r}"].value is not None for r in range(6, 11)), "Five negative rows visible")
    check(values["Model_Cost"].column_dimensions["A"].width >= 28, "Model label width")
    check(values["Executive_Summary"]["F7"].alignment.wrap_text, "Executive assumption note wrapping")
    check(values["Methodology_Limits"]["A5"].alignment.wrap_text, "Methodology note wrapping")
    print("PASS: 11-sheet workbook cached formulas, row values, visible scope and layout checks")

    current_paths = [p for p in OUT.rglob("*") if p.is_file() and p.suffix.lower() in {".csv", ".json", ".md"} and p.name not in {"FINAL_QA_REVIEW.md", "missing_cost_sensitivity.csv", "pareto_frontier.csv"}]
    current_paths += [p for p in (D6 / "src").rglob("*") if p.is_file() and p.suffix.lower() in {".py", ".yaml"}]
    current_paths += [p for p in (D6 / "preflight").rglob("*") if p.is_file() and p.suffix.lower() in {".csv", ".md", ".py"} and str(p.relative_to(D6)).replace("\\", "/") not in LEGACY]
    for path in current_paths:
        content = path.read_text(encoding="utf-8-sig").lower()
        check(not any(term in content for term in STALE), f"Active stale string in {path.relative_to(D6)}")
    legacy_present = [relative for relative in LEGACY if (D6 / relative).exists()]
    print(f"PASS: active generated/source stale-string sweep ({len(current_paths)} files)")
    print(f"BLOCKER: {len(legacy_present)} superseded Llama-era audit/output files retained after automatic deletion review rejection")
    for relative in legacy_present:
        print(f"  {relative}")
    print(f"SECOND_PASS_CALCULATIONS_PASS: {checks} assertions; RELEASE_BLOCKERS={len(legacy_present)}")


if __name__ == "__main__":
    main()
